from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass
from typing import Any


MAX_UPLOAD_BYTES = 12 * 1024 * 1024


@dataclass
class DocumentPayload:
    filename: str
    mime_type: str
    raw: bytes
    text: str


def decode_payload(filename: str, mime_type: str, file_b64: str) -> DocumentPayload:
    try:
        raw = base64.b64decode(file_b64, validate=True)
    except Exception as exc:
        raise ValueError("file_base64 không hợp lệ") from exc
    if not raw:
        raise ValueError("File rỗng")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError(f"File vượt quá {MAX_UPLOAD_BYTES // (1024 * 1024)} MB")
    text = extract_text(filename, mime_type, raw)
    return DocumentPayload(filename=filename or "document", mime_type=mime_type or "application/octet-stream", raw=raw, text=text)


def extract_text(filename: str, mime_type: str, raw: bytes) -> str:
    name = (filename or "").lower()
    mime = (mime_type or "").lower()
    if name.endswith(".txt") or mime.startswith("text/"):
        return raw.decode("utf-8", errors="replace").strip()
    if name.endswith(".docx") or "wordprocessingml" in mime:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("Thiếu python-docx trên server") from exc
        doc = Document(io.BytesIO(raw))
        parts: list[str] = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if vals:
                    parts.append(" | ".join(vals))
        return "\n".join(parts).strip()
    if name.endswith(".pdf") or mime == "application/pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return ""
        try:
            reader = PdfReader(io.BytesIO(raw))
            return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
        except Exception:
            return ""
    return ""


DATE_RE = re.compile(r"\b(?:ngày\s+)?([0-3]?\d)[/\-.]([01]?\d)[/\-.](20\d{2})\b", re.I)
DOC_NO_RE = re.compile(r"\b(?:(?:Số\s*:\s*)|(?:(?:Công văn|Văn bản|Tờ trình|Báo cáo|Kế hoạch)\s+(?:số\s+)?))[0-9]{1,6}\s*/\s*[A-ZĐ][A-ZĐ0-9.\-]{1,30}\b", re.I)
MONEY_RE = re.compile(r"\b\d{1,3}(?:[.,]\d{3})+(?:\s*(?:đồng|VNĐ|VND))?\b", re.I)
PERCENT_RE = re.compile(r"\b\d+(?:[.,]\d+)?\s*%")
LEGAL_RE = re.compile(
    r"\b(?:Nghị định|Thông tư|Quyết định|Nghị quyết|Chỉ thị|Hướng dẫn|Quy định|Kế hoạch)\s+"
    r"(?:số\s+)?[0-9A-ZĐ./\-]+(?:/[A-ZĐ0-9.\-]+)?",
    re.I,
)
DEADLINE_RE = re.compile(r"(?:trước|chậm nhất|hạn(?:\s+gửi)?|hoàn thành)\s+(?:ngày\s+)?[^\n.;]{0,45}", re.I)


def extract_protected_facts(text: str) -> list[dict[str, Any]]:
    facts: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    rules = [
        ("document_number", DOC_NO_RE),
        ("date", DATE_RE),
        ("money", MONEY_RE),
        ("percentage", PERCENT_RE),
        ("legal_reference", LEGAL_RE),
        ("deadline", DEADLINE_RE),
    ]
    for kind, pattern in rules:
        for m in pattern.finditer(text or ""):
            value = m.group(0).strip()
            key = (kind, value.lower())
            if value and key not in seen:
                facts.append({"type": kind, "value": value, "editable": False})
                seen.add(key)
    return facts[:120]


def compact_text(text: str, max_chars: int = 30000) -> str:
    text = re.sub(r"[ \t]+", " ", text or "")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) <= max_chars:
        return text
    half = max_chars // 2
    return text[:half] + "\n\n[...đã rút gọn...]\n\n" + text[-half:]
