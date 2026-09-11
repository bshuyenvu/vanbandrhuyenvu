from __future__ import annotations

import io
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Times New Roman"
PARTY_TITLE = "ĐẢNG CỘNG SẢN VIỆT NAM"


def _font_run(run, size: float, *, bold: bool = False, italic: bool = False, underline: bool = False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.underline = underline
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)
    return run


def _remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is not None:
        tbl_pr.remove(borders)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _set_cell_width(cell, width_cm: float):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def _tight(p, *, before: float = 0, after: float = 0, exact_line: float | None = None):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if exact_line is not None:
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(exact_line)


def _append_page_field(paragraph):
    run = paragraph.add_run()
    _font_run(run, 13)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def setup_party_page(doc: Document):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(14)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)

    header = section.header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p)
    _append_page_field(p)
    return section


def add_party_header(doc: Document, data: dict[str, Any]):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _remove_table_borders(table)
    left, right = table.rows[0].cells
    _set_cell_width(left, 7.7)
    _set_cell_width(right, 8.8)
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    parent = str(data.get("parent_agency") or "").strip().upper()
    agency = str(data.get("agency") or "TỔ CHỨC, CƠ QUAN ĐẢNG").strip().upper()

    left.text = ""
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p)
    if parent:
        _font_run(p.add_run(parent), 14)
        p = left.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p)
    _font_run(p.add_run(agency), 14, bold=True)

    p = left.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p, after=2)
    _font_run(p.add_run("*"), 14)

    number = str(data.get("number") or "…").strip()
    symbol = str(data.get("symbol") or "CV/…").strip()
    p = left.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p)
    _font_run(p.add_run(f"Số {number}-{symbol}"), 14)

    subject = str(data.get("subject") or "").strip()
    if subject:
        p = left.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p, before=1)
        _font_run(p.add_run(f"V/v {subject}"), 12, italic=True)

    right.text = ""
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p)
    _font_run(p.add_run(PARTY_TITLE), 15, bold=True)

    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p, before=0, after=4)
    _font_run(p.add_run("________________________"), 10)

    place = str(data.get("place") or "").strip()
    day = str(data.get("day") or "…").strip()
    month = str(data.get("month") or "…").strip()
    year = str(data.get("year") or "…").strip()
    prefix = f"{place}, " if place else ""
    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p)
    _font_run(p.add_run(f"{prefix}ngày {day} tháng {month} năm {year}"), 14, italic=True)
    return table


def add_party_recipient(doc: Document, recipient: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p, before=12, after=6, exact_line=20)
    _font_run(p.add_run("Kính gửi: "), 14, italic=True)
    _font_run(p.add_run(recipient or "[CẦN BỔ SUNG: cơ quan nhận]"), 14)


def add_party_body(doc: Document, paragraphs: Iterable[str]):
    for text in paragraphs:
        value = str(text or "").strip()
        if not value:
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.0)
        _tight(p, after=6, exact_line=20)
        _font_run(p.add_run(value), 14)


def add_party_signature(doc: Document, data: dict[str, Any]):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _remove_table_borders(table)
    left, right = table.rows[0].cells
    _set_cell_width(left, 7.3)
    _set_cell_width(right, 9.2)
    left.vertical_alignment = right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    left.text = ""
    p = left.paragraphs[0]
    _tight(p, before=12)
    _font_run(p.add_run("Nơi nhận:"), 14, underline=True)
    recipients = data.get("recipients") or ["- Như trên;"]
    for item in recipients:
        p = left.add_paragraph()
        _tight(p)
        _font_run(p.add_run(str(item)), 12)

    right.text = ""
    sign_authority = str(data.get("sign_authority") or "").strip().upper()
    signer_title = str(data.get("signer_title") or "").strip().upper()
    signer_name = str(data.get("signer_name") or "").strip()

    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p, before=12)
    if sign_authority:
        _font_run(p.add_run(sign_authority), 14, bold=True)
        p = right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tight(p)
    if signer_title:
        _font_run(p.add_run(signer_title), 14)
    else:
        _font_run(p.add_run("[CẦN BỔ SUNG: CHỨC VỤ NGƯỜI KÝ]"), 14)

    for _ in range(4):
        p = right.add_paragraph()
        _tight(p, exact_line=18)

    p = right.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tight(p)
    _font_run(p.add_run(signer_name or "[CẦN BỔ SUNG: Họ tên]"), 14, bold=True)
    return table


def build_party_reply(data: dict[str, Any]) -> bytes:
    """Build a Party reply *công văn* using deterministic 05-HD/VPTW layout rules."""
    doc = Document()
    setup_party_page(doc)
    add_party_header(doc, data)
    add_party_recipient(doc, str(data.get("recipient") or ""))
    paragraphs = data.get("paragraphs") or []
    if isinstance(paragraphs, str):
        paragraphs = [p.strip() for p in paragraphs.split("\n\n") if p.strip()]
    normalized = []
    for p in paragraphs:
        if isinstance(p, dict):
            p = p.get("text", "")
        if str(p).strip():
            normalized.append(str(p).strip())
    add_party_body(doc, normalized)
    add_party_signature(doc, data)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
