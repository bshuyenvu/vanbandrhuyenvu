from __future__ import annotations

import io
import os
import tempfile
import unittest
from pathlib import Path

from docx import Document

from webapp.government_docx import build_government_named_document
from webapp.saas import store
from webapp.saas.catalog import MODEL_DEFAULTS, quote_usage
from webapp.saas.document_types import get_document_type, list_document_types
from webapp.saas.model_manager import choose_model
from webapp.saas.rbac import has_permission
from webapp.saas.security import create_session, hash_password, parse_session, verify_password


class SecurityTests(unittest.TestCase):
    def setUp(self):
        self.old = os.environ.get("VBHC_SESSION_SECRET")
        os.environ["VBHC_SESSION_SECRET"] = "test-secret-0123456789-abcdefghijklmnopqrstuvwxyz"

    def tearDown(self):
        if self.old is None:
            os.environ.pop("VBHC_SESSION_SECRET", None)
        else:
            os.environ["VBHC_SESSION_SECRET"] = self.old

    def test_password_and_session(self):
        hashed = hash_password("MatKhau123!")
        self.assertTrue(verify_password("MatKhau123!", hashed))
        self.assertFalse(verify_password("sai-mat-khau", hashed))
        token = create_session({"uid": "usr_demo"}, ttl_seconds=60)
        self.assertEqual(parse_session(token)["uid"], "usr_demo")
        self.assertIsNone(parse_session(token + "x"))


class PolicyTests(unittest.TestCase):
    def test_rbac(self):
        self.assertTrue(has_permission("records_clerk", "document.issue"))
        self.assertTrue(has_permission("organization_owner", "document.approve"))
        self.assertFalse(has_permission("viewer", "document.issue"))

    def test_document_registry(self):
        self.assertTrue(get_document_type("gov_bao_cao").enabled_export)
        self.assertTrue(get_document_type("party_cong_van").enabled_export)
        self.assertFalse(get_document_type("party_bao_cao").enabled_export)
        outgoing = list_document_types(direction="outgoing")
        self.assertTrue(any(x["id"] == "gov_quyet_dinh" for x in outgoing))

    def test_lowest_cost_model_in_tier(self):
        choice = choose_model(task_type="draft", plan_id="personal", registry=MODEL_DEFAULTS)
        self.assertEqual(choice.model_id, "gpt-luna")
        with self.assertRaises(ValueError):
            choose_model(task_type="draft", plan_id="organization", data_policy="restricted", registry=MODEL_DEFAULTS)

    def test_credit_quote(self):
        q = quote_usage(input_tokens=1_000_000, output_tokens=1_000_000, model=MODEL_DEFAULTS["gpt-luna"])
        self.assertGreater(q.charged_credits, 0)
        self.assertGreater(q.provider_cost_usd, 0)


class WalletTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.old_db = store.DB_PATH
        store.DB_PATH = Path(self.tmp.name) / "test.db"
        store.init_db()

    def tearDown(self):
        store.DB_PATH = self.old_db
        self.tmp.cleanup()

    def test_usage_deducts_wallet_and_writes_ledger(self):
        wallet = store.create_personal_wallet("usr_test", "personal")
        before = wallet["balance_credits"]
        result = store.charge_ai_usage(
            user_id="usr_test", organization_id=None, department_id=None,
            provider="openai", model_name="gpt-5.6-luna", task_type="draft",
            usage={"input_tokens": 5000, "cached_tokens": 1000, "output_tokens": 1200},
        )
        after = store.one("SELECT balance_credits FROM wallets WHERE id=?", (wallet["id"],))["balance_credits"]
        self.assertEqual(before - after, result["charged_credits"])
        self.assertEqual(len(store.all_rows("SELECT * FROM ai_usage")), 1)
        self.assertGreaterEqual(len(store.all_rows("SELECT * FROM credit_ledger")), 2)


class RendererTests(unittest.TestCase):
    def test_government_report_renderer(self):
        raw = build_government_named_document({
            "document_type": "gov_bao_cao",
            "parent_agency": "UBND TỈNH",
            "agency": "ĐƠN VỊ A",
            "agency_abbr": "DVA",
            "subject": "Kết quả thực hiện nhiệm vụ tháng 9 năm 2026",
            "place": "Vĩnh Long",
            "month": 9,
            "year": 2026,
            "paragraphs": ["Đơn vị báo cáo kết quả thực hiện nhiệm vụ như sau."],
            "signer_title": "GIÁM ĐỐC",
            "signer_name": "Nguyễn Văn A",
        })
        self.assertTrue(raw.startswith(b"PK"))
        doc = Document(io.BytesIO(raw))
        text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
        self.assertIn("BÁO CÁO", text)
        self.assertIn("BC-DVA", text)
        self.assertIn("Kết quả thực hiện nhiệm vụ", text)


if __name__ == "__main__":
    unittest.main()
