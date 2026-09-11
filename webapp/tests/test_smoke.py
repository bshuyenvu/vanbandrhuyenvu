from __future__ import annotations

import io
import unittest

from docx import Document

from webapp.document_intake import extract_protected_facts, extract_text


class IntakeTests(unittest.TestCase):
    def test_document_number_does_not_confuse_dates_or_legal_refs(self):
        text = (
            "Số: 1234/SYT-NVY ngày 08/09/2026, căn cứ Nghị định "
            "30/2020/NĐ-CP, chậm nhất ngày 15/09/2026."
        )
        facts = extract_protected_facts(text)
        numbers = [x["value"] for x in facts if x["type"] == "document_number"]
        self.assertEqual(numbers, ["Số: 1234/SYT-NVY"])

    def test_docx_paragraph_and_table_extraction(self):
        doc = Document()
        doc.add_paragraph("CÔNG VĂN SỐ 12/ABC")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Yêu cầu"
        table.cell(0, 1).text = "Báo cáo"
        out = io.BytesIO()
        doc.save(out)
        text = extract_text(
            "sample.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            out.getvalue(),
        )
        self.assertIn("CÔNG VĂN", text)
        self.assertIn("Yêu cầu | Báo cáo", text)


if __name__ == "__main__":
    unittest.main()
