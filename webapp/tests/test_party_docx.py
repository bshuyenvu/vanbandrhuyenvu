from __future__ import annotations

import io
import unittest

from docx import Document
from docx.shared import Cm

from webapp.party_docx import PARTY_TITLE, build_party_reply


class PartyDocxTests(unittest.TestCase):
    def test_party_reply_layout(self):
        raw = build_party_reply({
            "parent_agency": "ĐẢNG ỦY CẤP TRÊN",
            "agency": "ĐẢNG ỦY ĐƠN VỊ",
            "number": "12",
            "symbol": "CV/ĐU",
            "place": "Vĩnh Long",
            "day": "11",
            "month": "09",
            "year": "2026",
            "subject": "phúc đáp báo cáo số liệu",
            "recipient": "Cơ quan cấp trên",
            "paragraphs": ["Thực hiện văn bản nêu trên, đơn vị báo cáo như sau."],
            "signer_title": "PHÓ BÍ THƯ",
            "signer_name": "Nguyễn Văn A",
        })
        self.assertTrue(raw.startswith(b"PK"))
        doc = Document(io.BytesIO(raw))
        sec = doc.sections[0]
        self.assertAlmostEqual(sec.top_margin / Cm(1), 2.0, places=1)
        self.assertAlmostEqual(sec.bottom_margin / Cm(1), 2.0, places=1)
        self.assertAlmostEqual(sec.left_margin / Cm(1), 3.0, places=1)
        self.assertAlmostEqual(sec.right_margin / Cm(1), 1.5, places=1)
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for row in t.rows for c in row.cells]
        )
        self.assertIn(PARTY_TITLE, text)
        self.assertIn("Số 12-CV/ĐU", text)
        self.assertIn("Kính gửi", text)
        self.assertIn("Nơi nhận", text)
        self.assertIn("PHÓ BÍ THƯ", text)


if __name__ == "__main__":
    unittest.main()
