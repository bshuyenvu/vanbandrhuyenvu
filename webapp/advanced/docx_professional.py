from __future__ import annotations

import io
import re
import zipfile
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, instr, separate])
    if placeholder: paragraph.add_run(placeholder)
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def _font(run, name: str = "Times New Roman", size: float = 13) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def _section(section, spec: dict[str, Any]) -> None:
    section.top_margin = Cm(float(spec.get("top_cm", 2)))
    section.bottom_margin = Cm(float(spec.get("bottom_cm", 2)))
    section.left_margin = Cm(float(spec.get("left_cm", 3)))
    section.right_margin = Cm(float(spec.get("right_cm", 2)))


def _update_fields(doc: Document) -> None:
    settings = doc.settings._element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields"); settings.append(node)
    node.set(qn("w:val"), "true")


def _notes(text: str, offset: int) -> tuple[str, list[tuple[int,str]]]:
    found: list[tuple[int,str]] = []
    def repl(m):
        idx = offset + len(found) + 1; found.append((idx, m.group(1).strip())); return f"[[HVFN{idx}]]"
    return re.sub(r"\{fn:(.+?)\}", repl, text), found


def _xml_escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _inject_footnotes(raw: bytes, notes: list[tuple[int,str]]) -> bytes:
    if not notes: return raw
    src, out = io.BytesIO(raw), io.BytesIO()
    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        document_xml = zin.read("word/document.xml").decode("utf-8")
        for idx, _ in notes:
            marker = f"[[HVFN{idx}]]"
            repl = f'</w:t></w:r><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteReference w:id="{idx}"/></w:r><w:r><w:t>'
            document_xml = document_xml.replace(marker, repl)
        rels = zin.read("word/_rels/document.xml.rels").decode("utf-8")
        if "relationships/footnotes" not in rels:
            rels = rels.replace("</Relationships>", '<Relationship Id="rIdHVFootnotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
        cts = zin.read("[Content_Types].xml").decode("utf-8")
        if "/word/footnotes.xml" not in cts:
            cts = cts.replace("</Types>", '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
        body = ['<w:footnote w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>', '<w:footnote w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>']
        for idx, note in notes:
            body.append(f'<w:footnote w:id="{idx}"><w:p><w:r><w:t xml:space="preserve">{_xml_escape(note)}</w:t></w:r></w:p></w:footnote>')
        foot_xml = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="%s">%s</w:footnotes>' % (NS_W, "".join(body))
        for info in zin.infolist():
            if info.filename in {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml", "word/footnotes.xml"}: continue
            zout.writestr(info, zin.read(info.filename))
        zout.writestr("word/document.xml", document_xml); zout.writestr("word/_rels/document.xml.rels", rels)
        zout.writestr("[Content_Types].xml", cts); zout.writestr("word/footnotes.xml", foot_xml)
    return out.getvalue()


def build_professional_docx(spec: dict[str, Any]) -> bytes:
    doc = Document(); margins = spec.get("margins", {}); _section(doc.sections[0], margins); _update_fields(doc)
    font_name, font_size = spec.get("font", "Times New Roman"), float(spec.get("font_size", 13))
    normal = doc.styles["Normal"]; normal.font.name = font_name; normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name); normal.font.size = Pt(font_size)
    header = str(spec.get("header") or "")
    if header:
        p = doc.sections[0].header.paragraphs[0]; p.text = header; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp = doc.sections[0].footer.paragraphs[0]
    footer = str(spec.get("footer") or "")
    if footer: fp.add_run(footer + " — ")
    if spec.get("page_numbers", True): _field(fp, "PAGE", "1")
    if spec.get("title"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(spec["title"])); _font(r, font_name, 16); r.bold = True
    if spec.get("toc"):
        p = doc.add_paragraph(); r = p.add_run("MỤC LỤC"); _font(r, font_name, 14); r.bold = True
        _field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "Cập nhật mục lục trong Word")
    notes: list[tuple[int,str]] = []
    for block in spec.get("blocks", []):
        kind = block.get("type", "paragraph")
        if kind == "section_break":
            sec = doc.add_section(WD_SECTION.NEW_PAGE); _section(sec, block.get("margins", margins)); continue
        if kind == "heading":
            p = doc.add_paragraph(style=f"Heading {max(1,min(int(block.get('level',1)),9))}")
            r = p.add_run(str(block.get("text") or "")); _font(r, font_name, 14 if int(block.get("level",1)) == 1 else 13); r.bold = True; continue
        if kind == "table":
            rows = block.get("rows") or []
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); table.style = "Table Grid"
                for i,row in enumerate(rows):
                    for j,value in enumerate(row): table.cell(i,j).text = str(value)
            if block.get("caption"):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(str(block["caption"])).bold = True
            continue
        if kind == "caption":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(str(block.get("text") or "")).italic = True; continue
        text, new_notes = _notes(str(block.get("text") or ""), len(notes)); notes.extend(new_notes)
        p = doc.add_paragraph(); r = p.add_run(text); _font(r, font_name, font_size)
        p.paragraph_format.first_line_indent = Cm(float(block.get("first_line_cm", 1)))
        p.paragraph_format.line_spacing = float(block.get("line_spacing", 1.3))
    bio = io.BytesIO(); doc.save(bio); return _inject_footnotes(bio.getvalue(), notes)
